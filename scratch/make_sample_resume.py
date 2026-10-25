import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_resume(docx_path: str):
    doc = docx.Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.7)
        sec.bottom_margin = Inches(0.7)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # Style: Serif
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Georgia'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Header Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_name.add_run("ALEXANDER HAMILTON")
    run_name.font.name = 'Georgia'
    run_name.font.size = Pt(22)
    run_name.font.bold = True
    run_name.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D) # Navy

    # Subtitle / Contact
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Senior Software Architect  |  New York, NY  |  alex@example.com  |  ")
    r_sub.font.name = 'Georgia'
    r_sub.font.size = Pt(10)
    
    # Add a real hyperlink
    def add_hyperlink(paragraph, url, text, color="004B87"):
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = parse_xml(f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/><w:color w:val="{color}"/><w:u w:val="single"/></w:rPr><w:t>{text}</w:t></w:r></w:hyperlink>')
        paragraph._p.append(hyperlink)

    add_hyperlink(p_sub, "https://github.com/alex-hamilton", "portfolio.dev")

    def add_shaded_section(title: str):
        # Shaded section heading bar using a 1-cell table with shading
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.9)
        cell = tbl.cell(0, 0)
        # Background shading: Deep Slate Navy
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1B365D"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        # Margins inside cell
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="80" w:type="dxa"/><w:bottom w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        r = cp.add_run(title.upper())
        r.font.name = 'Georgia'
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # Crisp white on dark background

    # Summary
    add_shaded_section("Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(4)
    p_sum.paragraph_format.space_after = Pt(8)
    r_sum = p_sum.add_run(
        "Accomplished systems architect with 12+ years of experience engineering high-performance "
        "offline desktop engines, native compilers, and secure cross-platform applications. "
        "Recognized for visual fidelity standards and sub-second rendering pipelines."
    )
    r_sum.font.name = 'Georgia'

    # Experience
    add_shaded_section("Professional Experience")
    p_exp1 = doc.add_paragraph()
    p_exp1.paragraph_format.space_before = Pt(4)
    p_exp1.paragraph_format.space_after = Pt(2)
    r_role = p_exp1.add_run("Lead Systems Engineer  |  Apex Architecture Labs")
    r_role.font.bold = True
    r_role.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    r_date = p_exp1.add_run("          2020 – Present")
    r_date.font.italic = True

    # Bullet points
    bullets = [
        "Architected native document conversion engine achieving 99.98% visual fidelity across complex vector layouts.",
        "Engineered zero-loss font metric mapping, preventing layout reflow and font substitution degradation.",
        "Led offline-first desktop initiative, eliminating runtime network dependencies and reducing latency by 85%."
    ]
    for b in bullets:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(2)
        bp.paragraph_format.space_after = Pt(2)
        br = bp.add_run(b)
        br.font.name = 'Georgia'
        br.font.size = Pt(10)

    # Education & Links
    add_shaded_section("Education & Credentials")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(4)
    p_edu.paragraph_format.space_after = Pt(2)
    r_deg = p_edu.add_run("B.S. in Computer Science  |  Columbia University")
    r_deg.font.bold = True
    
    p_link = doc.add_paragraph()
    p_link.paragraph_format.space_before = Pt(2)
    p_link.add_run("Verified Credentials: ")
    add_hyperlink(p_link, "https://credentials.columbia.edu/alex", "https://credentials.columbia.edu/alex")

    doc.save(docx_path)
    print(f"Created {docx_path}")

if __name__ == "__main__":
    create_resume("test_samples/sample_resume.docx")
