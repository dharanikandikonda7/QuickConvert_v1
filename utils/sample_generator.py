import os
from pathlib import Path
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas
from reportlab.lib.pdfencrypt import StandardEncryption


def generate_sample_files(target_dir: Path) -> dict[str, Path]:
    """Generates a complete suite of valid and edge-case test files."""
    target_dir.mkdir(parents=True, exist_ok=True)
    generated = {}

    # 1. Multi-page PDF
    pdf_path = target_dir / "sample_multipage.pdf"
    c = canvas.Canvas(str(pdf_path), pagesize=(400, 300))
    # Page 1
    c.setStrokeColorRGB(0.2, 0.4, 0.8)
    c.setFillColorRGB(0.9, 0.95, 1.0)
    c.rect(20, 20, 360, 260, fill=1, stroke=1)
    c.setFillColorRGB(0.1, 0.1, 0.3)
    c.setFont("Helvetica", 14)
    c.drawString(40, 240, "QuickConvert Test Document - Page 1")
    c.setFont("Helvetica", 10)
    c.setFillColorRGB(0, 0, 0)
    c.drawString(40, 200, "Testing PDF to PNG, JPEG, and TXT conversions.")
    c.showPage()
    # Page 2
    c.setStrokeColorRGB(0.8, 0.3, 0.2)
    c.setFillColorRGB(1.0, 0.95, 0.9)
    c.rect(20, 20, 360, 260, fill=1, stroke=1)
    c.setFillColorRGB(0.3, 0.1, 0.1)
    c.setFont("Helvetica", 14)
    c.drawString(40, 240, "QuickConvert Test Document - Page 2")
    c.setFont("Helvetica", 10)
    c.setFillColorRGB(0, 0, 0)
    c.drawString(40, 200, "This second page ensures multipage rendering works properly.")
    c.showPage()
    c.save()
    generated["pdf_multipage"] = pdf_path

    # 2. Encrypted PDF
    enc_path = target_dir / "sample_encrypted.pdf"
    enc = StandardEncryption("secret123", "owner123", canPrint=0)
    c_enc = canvas.Canvas(str(enc_path), pagesize=(300, 200), encrypt=enc)
    c_enc.drawString(30, 150, "Classified Content")
    c_enc.showPage()
    c_enc.save()
    generated["pdf_encrypted"] = enc_path

    # 3. Transparent PNG
    png_path = target_dir / "sample_transparent.png"
    im_png = Image.new("RGBA", (200, 200), (0, 0, 0, 0))
    draw = ImageDraw.Draw(im_png)
    draw.rounded_rectangle([20, 20, 180, 180], radius=30, fill=(99, 102, 241, 220), outline=(67, 56, 202, 255), width=3)
    draw.ellipse([60, 60, 140, 140], fill=(244, 114, 182, 240))
    im_png.save(png_path, "PNG")
    generated["png"] = png_path

    # 4. Standard JPEG
    jpg_path = target_dir / "sample_photo.jpg"
    im_jpg = Image.new("RGB", (250, 150), (240, 245, 250))
    draw_jpg = ImageDraw.Draw(im_jpg)
    draw_jpg.rectangle([10, 10, 240, 140], fill=(14, 165, 233), outline=(3, 105, 161), width=2)
    draw_jpg.text((30, 60), "QuickConvert JPEG Sample", fill=(255, 255, 255))
    im_jpg.save(jpg_path, "JPEG", quality=95)
    generated["jpeg"] = jpg_path

    # 5. WebP
    webp_path = target_dir / "sample_image.webp"
    im_jpg.save(webp_path, "WEBP", quality=90)
    generated["webp"] = webp_path

    # 6. BMP
    bmp_path = target_dir / "sample_graphic.bmp"
    im_jpg.save(bmp_path, "BMP")
    generated["bmp"] = bmp_path

    # 7. Multipage TIFF
    tiff_path = target_dir / "sample_multipage.tiff"
    t1 = Image.new("RGB", (180, 120), (220, 252, 231))
    d1 = ImageDraw.Draw(t1)
    d1.text((20, 50), "TIFF Page 1", fill=(22, 101, 52))

    t2 = Image.new("RGB", (180, 120), (254, 240, 138))
    d2 = ImageDraw.Draw(t2)
    d2.text((20, 50), "TIFF Page 2", fill=(133, 77, 14))

    t1.save(tiff_path, "TIFF", save_all=True, append_images=[t2])
    generated["tiff_multipage"] = tiff_path

    # 8. Multiframe GIF
    gif_path = target_dir / "sample_animated.gif"
    g1 = Image.new("RGB", (100, 100), (255, 228, 230))
    dg1 = ImageDraw.Draw(g1)
    dg1.ellipse([20, 20, 80, 80], fill=(225, 29, 72))

    g2 = Image.new("RGB", (100, 100), (224, 242, 254))
    dg2 = ImageDraw.Draw(g2)
    dg2.rectangle([20, 20, 80, 80], fill=(2, 132, 199))

    g1.save(gif_path, "GIF", save_all=True, append_images=[g2], duration=200, loop=0)
    generated["gif"] = gif_path

    # 9. Plain Unicode Text
    txt_path = target_dir / "sample_document.txt"
    sample_text = (
        "QuickConvert: Modern Windows 11 Document Converter\n"
        "==================================================\n\n"
        "Features:\n"
        "- Radial glassmorphism user interface\n"
        "- High fidelity PDF to image & text conversions\n"
        "- Unicode text support: Привет мир, こんにちは, café, naïve\n"
        "- Strict overwrite safety and clean architecture\n\n"
        "Paragraph 2:\n"
        "QuickConvert automatically derives non-colliding output filenames\n"
        "and handles multipage exports with sequential numbering.\n"
    )
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(sample_text)
    generated["txt"] = txt_path

    # 10. Clean HTML
    html_path = target_dir / "sample_page.html"
    sample_html = """<!DOCTYPE html>
<html>
<head><title>Test Article</title></head>
<body>
    <h1>Heading Level 1</h1>
    <p>This is a test paragraph containing <b>bold</b> and <i>italic</i> words.</p>
    <ul>
        <li>First item</li>
        <li>Second item</li>
    </ul>
</body>
</html>"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(sample_html)
    generated["html"] = html_path

    # 11. Empty File (0 bytes)
    empty_path = target_dir / "sample_empty.txt"
    empty_path.touch()
    generated["empty"] = empty_path

    # 12. Corrupt File
    corrupt_path = target_dir / "sample_corrupt.pdf"
    with open(corrupt_path, "wb") as f:
        f.write(b"%PDF-1.4\nJUNK_INVALID_DATA_NOT_A_REAL_PDF_EOF")
    generated["corrupt_pdf"] = corrupt_path

    # 13. Misleading Extension (JPEG content with .png extension)
    misleading_path = target_dir / "sample_misleading.png"
    im_jpg.save(misleading_path, "JPEG")
    generated["misleading_png"] = misleading_path

    # 14. Word Document (.docx)
    docx_path = target_dir / "sample_word.docx"
    import docx
    w_doc = docx.Document()
    w_doc.add_heading("QuickConvert Word Document", 0)
    w_doc.add_paragraph("This is a test Word document for QuickConvert conversions.")
    w_doc.add_heading("Features Overview", level=1)
    w_doc.add_paragraph("Supports PDF, Images, Text, and Word conversions offline.")
    w_doc.save(str(docx_path))
    generated["docx"] = docx_path

    return generated


if __name__ == "__main__":
    out = generate_sample_files(Path(__file__).parent.parent / "test_samples")
    print(f"Generated {len(out)} sample files in test_samples/")
